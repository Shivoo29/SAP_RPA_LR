"""
Excel Manager
=============
Handles Excel file operations for input and output.
"""

import logging
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from config import Config
from data.data_models import ProcessingResult, BatchProcessingReport


class ExcelManager:
    """Manages Excel file operations."""
    
    def __init__(self):
        """Initialize Excel manager."""
        self.logger = logging.getLogger(__name__)
        self.config = Config()
    
    def read_input_file(self, file_path: str) -> List[Dict]:
        """
        Read input Excel file with material numbers.
        
        Args:
            file_path: Path to Excel file
            
        Returns:
            List of dictionaries with material data
        """
        try:
            self.logger.info(f"Reading input file: {file_path}")
            
            # Read Excel file
            df = pd.read_excel(file_path)
            
            # Validate columns
            if df.empty:
                self.logger.error("Input file is empty")
                return []
            
            # Get material numbers from first column
            # Assuming format: Part Number | Plant (optional) | MRP Area (optional)
            materials = []
            
            for idx, row in df.iterrows():
                material_data = {
                    'material_number': str(row.iloc[0]).strip(),
                    'plant': str(row.iloc[1]).strip() if len(row) > 1 else None,
                    'mrp_area': str(row.iloc[2]).strip() if len(row) > 2 else None
                }
                
                # Skip empty rows
                if material_data['material_number'] and material_data['material_number'] != 'nan':
                    materials.append(material_data)
            
            self.logger.info(f"Read {len(materials)} materials from input file")
            return materials
            
        except Exception as e:
            self.logger.error(f"Error reading input file: {e}")
            return []
    
    def write_output_file(
        self,
        results: List[ProcessingResult],
        report: Optional[BatchProcessingReport] = None,
        output_path: Optional[str] = None
    ) -> str:
        """
        Write results to Excel file with multiple sheets.
        
        Args:
            results: List of ProcessingResult objects
            report: Optional batch processing report
            output_path: Optional custom output path
            
        Returns:
            Path to created file
        """
        try:
            # Generate filename if not provided
            if not output_path:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_dir = Path(self.config.OUTPUT_DIRECTORY)
                output_dir.mkdir(exist_ok=True)
                output_path = output_dir / f"{self.config.OUTPUT_FILE_PREFIX}_{timestamp}.xlsx"
            
            self.logger.info(f"Writing output file: {output_path}")
            
            # Create Excel writer
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                
                # Sheet 1: Main Results
                self._write_results_sheet(writer, results)
                
                # Sheet 2: Summary
                if report:
                    self._write_summary_sheet(writer, report)
                else:
                    self._write_summary_sheet_from_results(writer, results)
                
                # Sheet 3: Errors
                self._write_errors_sheet(writer, results)
            
            # Apply formatting
            self._apply_formatting(output_path)
            
            self.logger.info(f"Successfully created output file: {output_path}")
            return str(output_path)
            
        except Exception as e:
            self.logger.error(f"Error writing output file: {e}")
            raise
    
    def _write_results_sheet(self, writer, results: List[ProcessingResult]):
        """Write main results sheet."""
        # Convert results to dictionaries
        data = [result.to_dict() for result in results]
        
        # Create DataFrame
        df = pd.DataFrame(data)
        
        # Reorder columns for better readability
        primary_cols = ['material_number', 'success', 'scenario', 'plant_found']
        data_cols = [col for col in df.columns if col not in primary_cols + ['timestamp', 'processing_time', 'error_message']]
        final_cols = primary_cols + data_cols + ['processing_time', 'timestamp', 'error_message']
        
        # Filter existing columns
        final_cols = [col for col in final_cols if col in df.columns]
        df = df[final_cols]
        
        # Write to Excel
        sheet_name = self.config.EXCEL_SHEET_NAMES['results']
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        self.logger.debug(f"Wrote {len(df)} rows to results sheet")
    
    def _write_summary_sheet(self, writer, report: BatchProcessingReport):
        """Write summary sheet from report."""
        summary_data = {
            'Metric': [
                'Total Materials Processed',
                'Successful',
                'Failed',
                'Success Rate (%)',
                'Scenario 1 (MD04 MatRes Found)',
                'Scenario 2 (ERF Dashboard)',
                'Scenario 3 (ERF → KO03)',
                'Processing Start Time',
                'Processing End Time',
                'Total Processing Time (seconds)',
                'Average Time per Material (seconds)'
            ],
            'Value': [
                report.total_materials,
                report.successful,
                report.failed,
                round(report.success_rate, 2),
                report.scenario_1_count,
                report.scenario_2_count,
                report.scenario_3_count,
                report.start_time.strftime('%Y-%m-%d %H:%M:%S'),
                report.end_time.strftime('%Y-%m-%d %H:%M:%S') if report.end_time else 'N/A',
                round(report.processing_time, 2),
                round(report.processing_time / report.total_materials, 2) if report.total_materials > 0 else 0
            ]
        }
        
        df = pd.DataFrame(summary_data)
        sheet_name = self.config.EXCEL_SHEET_NAMES['summary']
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        self.logger.debug("Wrote summary sheet")
    
    def _write_summary_sheet_from_results(self, writer, results: List[ProcessingResult]):
        """Write summary sheet directly from results."""
        total = len(results)
        successful = len([r for r in results if r.success])
        failed = total - successful
        
        # Count scenarios
        from data.data_models import ScenarioType
        scenario_1 = len([r for r in results if r.scenario == ScenarioType.MD04_MATRES_FOUND])
        scenario_2 = len([r for r in results if r.scenario == ScenarioType.ERF_DASHBOARD])
        scenario_3 = len([r for r in results if r.scenario == ScenarioType.ERF_TO_KO03])
        
        # Calculate times
        total_time = sum([r.processing_time for r in results])
        avg_time = total_time / total if total > 0 else 0
        
        summary_data = {
            'Metric': [
                'Total Materials Processed',
                'Successful',
                'Failed',
                'Success Rate (%)',
                'Scenario 1 (MD04 MatRes Found)',
                'Scenario 2 (ERF Dashboard)',
                'Scenario 3 (ERF → KO03)',
                'Total Processing Time (seconds)',
                'Average Time per Material (seconds)'
            ],
            'Value': [
                total,
                successful,
                failed,
                round((successful / total * 100), 2) if total > 0 else 0,
                scenario_1,
                scenario_2,
                scenario_3,
                round(total_time, 2),
                round(avg_time, 2)
            ]
        }
        
        df = pd.DataFrame(summary_data)
        sheet_name = self.config.EXCEL_SHEET_NAMES['summary']
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        self.logger.debug("Wrote summary sheet from results")
    
    def _write_errors_sheet(self, writer, results: List[ProcessingResult]):
        """Write errors sheet with failed materials."""
        # Filter failed results
        failed_results = [r for r in results if not r.success]
        
        if not failed_results:
            # Create empty sheet with message
            df = pd.DataFrame({'Message': ['No errors - all materials processed successfully!']})
        else:
            data = [{
                'material_number': r.material_number,
                'scenario_attempted': r.scenario.value,
                'error_message': r.error_message,
                'timestamp': r.timestamp.strftime('%Y-%m-%d %H:%M:%S')
            } for r in failed_results]
            
            df = pd.DataFrame(data)
        
        sheet_name = self.config.EXCEL_SHEET_NAMES['errors']
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        
        self.logger.debug(f"Wrote {len(failed_results)} errors to errors sheet")
    
    def _apply_formatting(self, file_path: str):
        """Apply formatting to Excel file."""
        try:
            wb = load_workbook(file_path)
            
            # Format each sheet
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Format headers
                header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
                header_font = Font(bold=True, color='FFFFFF')
                
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                
                # Auto-adjust column widths
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    
                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width
            
            # Save formatted workbook
            wb.save(file_path)
            
            self.logger.debug("Applied formatting to Excel file")
            
        except Exception as e:
            self.logger.warning(f"Could not apply formatting: {e}")
